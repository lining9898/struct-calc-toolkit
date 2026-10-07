from pathlib import Path
from functools import partial
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from threading import Thread
from playwright.sync_api import sync_playwright
from docx import Document
import zipfile,json,os
from lxml import etree
root=Path(__file__).resolve().parents[1];out=Path(os.environ.get('WORD_SAMPLES_DIR','/tmp/struct-calc-word-samples'));out.mkdir(parents=True,exist_ok=True)
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*a):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(root)));Thread(target=server.serve_forever,daemon=True).start()
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);page=browser.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)));page.on('dialog',lambda d:(print('DIALOG',d.message,flush=True),d.dismiss()));page.on('console',lambda m:print('CONSOLE',m.text,flush=True) if m.type=='error' else None)
 page.goto('http://127.0.0.1:'+str(server.server_port));page.wait_for_timeout(100)
 ids=page.evaluate('Object.keys(TOOLS)');results=[]
 for id in ids:
  page.evaluate('(id)=>goTool(id)',id);page.wait_for_timeout(120)
  with page.expect_download(timeout=15000) as dl:page.evaluate('exportBook()')
  file=out/(id+'.docx');dl.value.save_as(file)
  assert dl.value.suggested_filename.endswith('.docx'),id
  d=Document(file);assert len(d.paragraphs)>0,id
  with zipfile.ZipFile(file) as z:
   xml=z.read('word/document.xml').decode();assert 'w:pgSz' in xml and 'w:tblLayout' in xml,id
   assert 'w:footerReference' in xml,id
   assert 'html' not in z.namelist(),id
   ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
   tree=etree.fromstring(xml.encode());styles=etree.fromstring(z.read('word/styles.xml'))
   assert tree.find('.//w:pgSz',ns).get('{'+ns['w']+'}w')=='11906',id
   assert tree.find('.//w:pgSz',ns).get('{'+ns['w']+'}h')=='16838',id
   assert styles.find('.//w:rFonts',ns).get('{'+ns['w']+'}eastAsia')=='宋体',id
   for table in tree.findall('.//w:tbl',ns):
    assert sum(int(c.get('{'+ns['w']+'}w')) for c in table.findall('w:tblGrid/w:gridCol',ns))==9638,id
  assert 'w:vertAlign' in xml or not page.locator('.proc-body sub').count(),id
  results.append({'tool':id,'paragraphs':len(d.paragraphs),'tables':len(d.tables),'bytes':file.stat().st_size})
  print('PASS',id,flush=True)
 assert not errors,errors
 browser.close()
server.shutdown();(out/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));print('ALL',len(results),'native DOCX exports passed')
